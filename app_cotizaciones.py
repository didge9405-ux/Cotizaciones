import streamlit as st
from google import genai
from docx import Document
import io
import re

# Configuración de la página web
st.set_page_config(page_title="Sistema de Cotizaciones YAVIS", page_icon="📝", layout="centered")

st.title("📝 Generador de Cotizaciones YAVIS")
st.write("Escribe o dicta los datos de la cotización en lenguaje natural. La IA llenará la plantilla oficial automáticamente.")

# Campo de texto grande para IA
prompt_usuario = st.text_area(
    "Describe la cotización (Cliente, productos, cantidades, precios, etc.):",
    placeholder="Ejemplo: Cotización para Juan Pérez de 5 cámaras de seguridad a 1200 pesos cada una, con entrega en 3 días."
)

if st.button("🚀 Generar Cotización", type="primary"):
    if not prompt_usuario.strip():
        st.warning("Por favor ingresa o dicta los detalles de la cotización.")
    else:
        with st.spinner("Procesando con IA y generando documento..."):
            try:
                # Inicializar el cliente de Gemini (asegúrate de configurar tu API Key en los Secrets de Streamlit)
                # O puedes ingresar tu API key aquí temporalmente si prefieres:
                client = genai.Client(api_key=st.secrets.get("GEMINI_API_KEY", ""))

                # Prompt para extraer los datos estructurados con Gemini 2.5 Flash
                prompt_sistema = f"""
                Extrae de la siguiente descripción los datos para una cotización de YAVIS.
                Devuélvelos estrictamente en formato clave: valor, separados por saltos de línea.
                Las claves necesarias son:
                CLIENTE: [Nombre del cliente]
                FECHA: [Fecha actual o indicada]
                ITEMS: [Lista detallada de productos/servicios con cantidad, descripción y precio]
                TOTAL: [Monto total calculado]

                Texto del usuario: {prompt_usuario}
                """

                response = client.models.generate_content(
                    model='gemini-2.5-flash',
                    contents=prompt_sistema
                )
                
                texto_ia = response.text

                # Cargar la plantilla base de Word
                doc = Document("plantilla_base.docx")

                # Función para reemplazar etiquetas en los párrafos y tablas preservando formatos
                def reemplazar_en_parrafo(parrafo, datos_dict):
                    for clave, valor in datos_dict.items():
                        marcador = f"{{{{{clave}}}}}"
                        if marcador in parrafo.text:
                            for run in parrafo.runs:
                                if marcador in run.text:
                                    run.text = run.text.replace(marcador, str(valor))

                # Extraer pares clave-valor sencillos de la respuesta de la IA
                datos_dict = {}
                for linea in texto_ia.split("\n"):
                    if ":" in linea:
                        partes = linea.split(":", 1)
                        datos_dict[partes[0].strip().upper()] = partes[1].strip()

                # Reemplazar en párrafos del documento
                for p in doc.paragraphs:
                    reemplazar_en_parrafo(p, datos_dict)

                # Reemplazar en tablas del documento
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                reemplazar_en_parrafo(p, datos_dict)

                # Guardar el documento en memoria
                buffer = io.BytesIO()
                doc.save(buffer)
                buffer.seek(0)

                st.success("¡Cotización generada con éxito!")

                # Botón de descarga directa para el celular o PC
                st.download_button(
                    label="📥 Descargar Cotización en Word",
                    data=buffer,
                    file_name="Cotizacion_YAVIS.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

            except Exception as e:
                st.error(f"Ocurrió un error al procesar la cotización: {e}")
